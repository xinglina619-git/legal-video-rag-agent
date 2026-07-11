from html import escape
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def render_report_pdf(markdown: str, title: str = "Legal Material Review Draft") -> bytes:
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "CJKBody",
        parent=styles["BodyText"],
        fontName="STSong-Light",
        fontSize=9.5,
        leading=15,
        spaceAfter=6,
    )
    heading = ParagraphStyle(
        "CJKHeading",
        parent=styles["Heading2"],
        fontName="STSong-Light",
        fontSize=14,
        leading=20,
        spaceBefore=10,
        spaceAfter=6,
    )
    stream = BytesIO()
    document = SimpleDocTemplate(
        stream,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=title,
    )
    story = [Paragraph(escape(title), heading), Spacer(1, 4)]
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            story.append(Spacer(1, 4))
            continue
        if line.startswith("#"):
            story.append(Paragraph(escape(line.lstrip("#").strip()), heading))
            continue
        if line.startswith(("- ", "* ")):
            line = "• " + line[2:]
        story.append(Paragraph(escape(line).replace("\n", "<br/>"), base))
    document.build(story)
    return stream.getvalue()


def render_report_docx(markdown: str, title: str = "法律纠纷材料分析报告") -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt
    except ImportError:
        return _render_basic_docx(markdown, title)
    stream = BytesIO()
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    document.add_heading(title, level=0)
    for raw_line in markdown.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith("#"):
            level = min(len(line) - len(line.lstrip("#")), 3)
            document.add_heading(line.lstrip("#").strip(), level=level)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:].strip(), style="List Bullet")
        else:
            document.add_paragraph(line)
    document.save(stream)
    return stream.getvalue()


def _render_basic_docx(markdown: str, title: str) -> bytes:
    """Standards-compliant minimal DOCX fallback for lean demo environments."""
    paragraphs = [title, *[line.strip("#* -") for line in markdown.splitlines() if line.strip()]]
    body = "".join(
        f"<w:p><w:r><w:t xml:space=\"preserve\">{escape(line)}</w:t></w:r></w:p>"
        for line in paragraphs
    )
    content_types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>"""
    relationships = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>"""
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body>{body}<w:sectPr/></w:body></w:document>"""
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("_rels/.rels", relationships)
        archive.writestr("word/document.xml", document_xml)
    return stream.getvalue()
